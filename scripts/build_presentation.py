"""Construit le support de soutenance (livrable de présentation).

Génère un diaporama PowerPoint résumant le projet pour la session de bilan avec le
mentor : contexte, architecture du pipeline, déroulé des cinq étapes, résultats
réels et bilan. Construit programmatiquement avec ``python-pptx`` pour rester
reproductible.

Usage :
    uv run python scripts/build_presentation.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
REPORTS_DIR = ROOT / "reports"
SCHEMA_PNG = ROOT / "docs" / "schema_donnees.png"

# Palette « bleu nuit / cyan » sobre.
BLEU_NUIT = RGBColor(0x12, 0x1E, 0x33)
CYAN = RGBColor(0x2E, 0xC4, 0xD6)
BLANC = RGBColor(0xF5, 0xF7, 0xFA)
GRIS = RGBColor(0x5A, 0x6A, 0x7A)

LARGEUR = Inches(13.333)
HAUTEUR = Inches(7.5)


def _fond(slide, couleur: RGBColor) -> None:
    """Applique une couleur de fond pleine à une diapositive."""
    fond = slide.background.fill
    fond.solid()
    fond.fore_color.rgb = couleur


def _zone_texte(slide, gauche, haut, largeur, hauteur):
    """Ajoute une zone de texte et renvoie son cadre."""
    boite = slide.shapes.add_textbox(gauche, haut, largeur, hauteur)
    return boite.text_frame


def slide_titre(prs: Presentation, titre: str, sous_titre: str) -> None:
    """Diapositive de titre (fond bleu nuit)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fond(slide, BLEU_NUIT)
    tf = _zone_texte(slide, Inches(1), Inches(2.4), Inches(11.3), Inches(2))
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = titre
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = BLANC
    p2 = tf.add_paragraph()
    run2 = p2.add_run()
    run2.text = sous_titre
    run2.font.size = Pt(22)
    run2.font.color.rgb = CYAN


def slide_contenu(prs: Presentation, titre: str, points: list[str]) -> None:
    """Diapositive titre + liste à puces."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fond(slide, BLANC)

    tf_titre = _zone_texte(slide, Inches(0.7), Inches(0.4), Inches(12), Inches(1))
    run = tf_titre.paragraphs[0].add_run()
    run.text = titre
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = BLEU_NUIT

    # Filet d'accent.
    ligne = slide.shapes.add_shape(1, Inches(0.7), Inches(1.35), Inches(3), Pt(3))
    ligne.fill.solid()
    ligne.fill.fore_color.rgb = CYAN
    ligne.line.fill.background()

    tf = _zone_texte(slide, Inches(0.8), Inches(1.7), Inches(11.8), Inches(5.2))
    for i, point in enumerate(points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = point
        run.font.size = Pt(18)
        run.font.color.rgb = BLEU_NUIT
        p.space_after = Pt(10)


def slide_image(prs: Presentation, titre: str, image: Path, legende: str) -> None:
    """Diapositive titre + image centrée."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fond(slide, BLANC)
    tf_titre = _zone_texte(slide, Inches(0.7), Inches(0.4), Inches(12), Inches(1))
    run = tf_titre.paragraphs[0].add_run()
    run.text = titre
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = BLEU_NUIT

    if image.exists():
        slide.shapes.add_picture(str(image), Inches(3.4), Inches(1.5), height=Inches(5))
    tf = _zone_texte(slide, Inches(0.8), Inches(6.7), Inches(11.8), Inches(0.6))
    run = tf.paragraphs[0].add_run()
    run.text = legende
    run.font.size = Pt(14)
    run.font.italic = True
    run.font.color.rgb = GRIS


def main() -> None:
    """Construit le diaporama et l'enregistre dans reports/."""
    prs = Presentation()
    prs.slide_width = LARGEUR
    prs.slide_height = HAUTEUR

    slide_titre(
        prs,
        "CheckItAI — Extraction de données multimodales",
        "Projet 12 · Pipeline ETL pour un détecteur de fake news · Benoit Girard",
    )
    slide_contenu(
        prs,
        "Contexte & mission",
        [
            "• CheckItAI : start-up de détection automatique de désinformation.",
            "• Besoin : alimenter un détecteur multimodal (texte + image) en données fraîches et fiables.",
            "• Mission : concevoir un pipeline ETL automatisé — extraction, transformation, chargement.",
            "• Rôle : ingénieur data junior, en autonomie, du sourcing jusqu'au monitoring.",
        ],
    )
    slide_contenu(
        prs,
        "Architecture du pipeline",
        [
            "1. Extract — 3 sources officielles : flux RSS, API NewsData.io, dataset FakeNewsNet.",
            "2. Transform — nettoyage, validation image, normalisation vers un schéma unique.",
            "3. Load — chargement en base relationnelle (SQLite local / PostgreSQL-Supabase).",
            "4. Orchestration — DAG Apache Airflow planifié quotidiennement (Docker).",
            "5. Pilotage — tableau de bord KPI (Streamlit) + plan de monitoring.",
            "Stack : uv · ruff · pytest · pre-commit · Airflow · Streamlit · Docker.",
        ],
    )
    slide_contenu(
        prs,
        "Étape 1 — Qualifier les sources",
        [
            "• Flux RSS (The Guardian, BBC News, ABC News) : texte + image, officiel, sans clé.",
            "• API NewsData.io : actualité JSON multilingue avec image_url.",
            "• FakeNewsNet : la vérité terrain (real / fake) — apprentissage supervisé.",
            "• Canaux officiels privilégiés, pas de scraping ; distinction opinion vs désinformation.",
        ],
    )
    slide_contenu(
        prs,
        "Étapes 2 & 3 — Extraction & transformation",
        [
            "• Code modulaire : un connecteur par source, fonctions fetch_* isolées (try/except + logs).",
            "• Transformation en 3 temps : lecture → traitement → export.",
            "• Petites fonctions nommées : nettoie_texte(), valide_image(), extrait_domaine().",
            "• Mode multimodal strict : garantit le lien texte-image (champ has_image).",
            "• Paramètres configurables ; export Parquet/CSV ; pipeline reproductible et testé.",
        ],
    )
    slide_image(
        prs,
        "Étape 3 — Schéma conceptuel des données",
        SCHEMA_PNG,
        "Modèle conceptuel : PUBLICATION au cœur, lien texte-image garanti (1—1), label optionnel.",
    )
    slide_contenu(
        prs,
        "Étape 4 — Orchestration Airflow",
        [
            "• DAG checkitai_etl : 4 tâches distinctes extract → transform → load → metriques.",
            "• Réutilise directement les fonctions du package (mêmes que les scripts).",
            "• Planification quotidienne (@daily), reprises automatiques (retries).",
            "• Exécution locale via Docker ; chargement en base sécurisable (auth, rôles, chiffrage).",
        ],
    )
    slide_contenu(
        prs,
        "Étape 5 — KPI & monitoring",
        [
            "• KPI qualité : taux de validité, taux d'association texte-image, taux labellisé.",
            "• KPI performance : durée par étape, débit, appels API (coût).",
            "• Tableau de bord Streamlit lisible pour un public non technique.",
            "• Plan de monitoring : seuils d'alerte 🟢🟠🔴, gestion d'erreurs, fréquences, dérive (Evidently).",
        ],
    )
    slide_contenu(
        prs,
        "Démonstration — résultats réels",
        [
            "• 6 sources contributrices : The Guardian, BBC, ABC News, NewsData.io, PolitiFact, GossipCop.",
            "• 142 publications brutes collectées → 141 publications valides multimodales.",
            "• ~99 % de validité et 100 % d'association texte-image (mode strict).",
            "• Pipeline complet exécuté en ~2 secondes ; 1 appel API consommé.",
            "• Données chargées en base et visualisées dans le tableau de bord.",
        ],
    )
    slide_contenu(
        prs,
        "Bilan — difficultés & points forts",
        [
            "• Difficulté : extraire l'image des flux RSS (emplacements multiples) — fonction dédiée.",
            "• Difficulté : exécuter Airflow proprement en local — résolu via Docker Compose.",
            "• Point fort : architecture modulaire réutilisée à l'identique par scripts, notebooks et DAG.",
            "• Point fort : robustesse (sources isolées, logs, tests) et reproductibilité (uv, config).",
            "• Suite : enrichir les labels (Hugging Face), brancher Evidently, déployer sur Supabase.",
        ],
    )
    slide_titre(prs, "Merci !", "Questions & démonstration en direct du pipeline")

    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    sortie = REPORTS_DIR / "presentation.pptx"
    prs.save(sortie)
    print(
        f"[ok] présentation enregistrée : {sortie} ({len(prs.slides.__iter__.__self__._sldIdLst)} diapositives)"
    )


if __name__ == "__main__":
    main()
