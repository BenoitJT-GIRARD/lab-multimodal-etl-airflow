"""Construit le support de soutenance.

⚠️ **Script d'outillage.** Il ne fait pas partie du pipeline — voir
`scripts/outillage/README.md`.

Génère un diaporama PowerPoint résumant le projet pour la session de bilan :
contexte, architecture du pipeline, déroulé des étapes, résultats réels et bilan.
Le fichier produit reste modifiable dans PowerPoint.

Usage :
    uv run python scripts/outillage/build_presentation.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
REPORTS_DIR = ROOT / "reports"
SCHEMA_PNG = ROOT / "docs" / "schema_donnees.png"

# Palette « bleu nuit / cyan » sobre.
BLEU_NUIT = RGBColor(0x12, 0x1E, 0x33)
CYAN = RGBColor(0x2E, 0xC4, 0xD6)
BLANC = RGBColor(0xF5, 0xF7, 0xFA)
GRIS = RGBColor(0x5A, 0x6A, 0x7A)

LARGEUR = Inches(13.333)
HAUTEUR = Inches(7.5)


def _fond(slide, couleur: RGBColor) -> None:
    """Applique une couleur de fond pleine à une diapositive."""
    fond = slide.background.fill
    fond.solid()
    fond.fore_color.rgb = couleur


def _zone_texte(slide, gauche, haut, largeur, hauteur):
    """Ajoute une zone de texte et renvoie son cadre."""
    boite = slide.shapes.add_textbox(gauche, haut, largeur, hauteur)
    return boite.text_frame


def slide_titre(prs: Presentation, titre: str, sous_titre: str) -> None:
    """Diapositive de titre (fond bleu nuit)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fond(slide, BLEU_NUIT)
    tf = _zone_texte(slide, Inches(1), Inches(2.4), Inches(11.3), Inches(2))
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = titre
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = BLANC
    p2 = tf.add_paragraph()
    run2 = p2.add_run()
    run2.text = sous_titre
    run2.font.size = Pt(22)
    run2.font.color.rgb = CYAN


def slide_contenu(prs: Presentation, titre: str, points: list[str]) -> None:
    """Diapositive titre + liste à puces."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fond(slide, BLANC)

    tf_titre = _zone_texte(slide, Inches(0.7), Inches(0.4), Inches(12), Inches(1))
    run = tf_titre.paragraphs[0].add_run()
    run.text = titre
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = BLEU_NUIT

    # Filet d'accent.
    ligne = slide.shapes.add_shape(1, Inches(0.7), Inches(1.35), Inches(3), Pt(3))
    ligne.fill.solid()
    ligne.fill.fore_color.rgb = CYAN
    ligne.line.fill.background()

    tf = _zone_texte(slide, Inches(0.8), Inches(1.7), Inches(11.8), Inches(5.2))
    for i, point in enumerate(points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = point
        run.font.size = Pt(18)
        run.font.color.rgb = BLEU_NUIT
        p.space_after = Pt(10)


def slide_image(prs: Presentation, titre: str, image: Path, legende: str) -> None:
    """Diapositive titre + image centrée."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fond(slide, BLANC)
    tf_titre = _zone_texte(slide, Inches(0.7), Inches(0.4), Inches(12), Inches(1))
    run = tf_titre.paragraphs[0].add_run()
    run.text = titre
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = BLEU_NUIT

    if image.exists():
        slide.shapes.add_picture(str(image), Inches(3.4), Inches(1.5), height=Inches(5))
    tf = _zone_texte(slide, Inches(0.8), Inches(6.7), Inches(11.8), Inches(0.6))
    run = tf.paragraphs[0].add_run()
    run.text = legende
    run.font.size = Pt(14)
    run.font.italic = True
    run.font.color.rgb = GRIS


def main() -> None:
    """Construit le diaporama et l'enregistre dans reports/."""
    prs = Presentation()
    prs.slide_width = LARGEUR
    prs.slide_height = HAUTEUR

    slide_titre(
        prs,
        "CheckItAI — Extraction de données multimodales",
        "Projet 12 · Pipeline ETL pour un détecteur de fake news · Benoit Girard",
    )
    slide_contenu(
        prs,
        "Le problème à résoudre",
        [
            "• CheckItAI développe un détecteur de désinformation multimodal : il lit le texte",
            "   d'une publication et regarde son image.",
            "• Sans alimentation continue, le jeu de données vieillit et le modèle se dégrade.",
            "• Mission : un pipeline ETL qui collecte, nettoie et stocke ces publications",
            "   chaque jour, sans intervention.",
        ],
    )
    slide_contenu(
        prs,
        "Quatre sources, quatre méthodes d'accès",
        [
            "• Flux RSS (The Guardian, BBC, ABC News) — volume et fraîcheur, sans clé.",
            "• API NewsData.io — actualité déjà normalisée, avec quota à surveiller.",
            "• FakeNewsNet, téléchargé depuis GitHub — vérité terrain académique.",
            "• Fakeddit, export Kaggle déposé localement — volume multimodal annoté.",
            "",
            "Aucun scraping : coût de maintenance permanent, conditions d'utilisation, et les",
            "flux officiels exposent déjà la même donnée.",
        ],
    )
    slide_contenu(
        prs,
        "Extraction — le texte ET l'image",
        [
            "• Un module par source, une fonction fetch_* ; chaque source isolée en try/except.",
            "• Les images sont téléchargées et ouvertes avec Pillow : une URL ne prouve rien.",
            "• Le JSON brut porte le texte et le chemin du fichier image, en relatif —",
            "   le jeu de données reste valable hors de la machine qui l'a produit.",
            "• FakeNewsNet ne contient aucune image : elle est retrouvée dans la balise",
            "   og:image publiée par l'éditeur. Rendement mesuré : environ une sur trois.",
        ],
    )
    slide_contenu(
        prs,
        "Transformation — lecture, traitement, export",
        [
            "• Petites fonctions nommées : nettoie_texte(), valide_image(), normalise_date(),",
            "   normalise_langue(), normalise_label().",
            "• valide_image() vérifie que le fichier est sur le disque, pas que l'URL existe :",
            "   c'est ce qui garantit l'association texte-image.",
            "• Paramètres configurables, journalisation à chaque étape, export en Parquet.",
        ],
    )
    slide_image(
        prs,
        "Le schéma conceptuel",
        SCHEMA_PNG,
        "PUBLICATION au cœur ; les relations 1—1 avec le texte et l'image expriment la règle "
        "du jeu de données. Diagramme généré depuis le code.",
    )
    slide_contenu(
        prs,
        "Chargement — un modèle relationnel qui s'enrichit",
        [
            "• Une table par entité du schéma, reliées par id et source_id.",
            "• Plus une table à plat, prête pour l'entraînement et le tableau de bord.",
            "• Chargement incrémental : seules les publications inconnues sont ajoutées.",
            "   Deuxième exécution : 179 collectées, 11 réellement nouvelles.",
            "• Sécurité : pas de secret en dur, rôle applicatif limité, noms de tables validés.",
        ],
    )
    slide_contenu(
        prs,
        "Orchestration — des tâches vraiment indépendantes",
        [
            "• DAG checkitai_etl : extraction → transformation → chargement → métriques → nettoyage.",
            "• Aucune donnée ne passe par XCom : chaque étape écrit son résultat dans",
            "   data/interim/ et l'étape suivante relit ce fichier.",
            "• N'importe quelle tâche se rejoue seule, et retombe sur le dernier artefact",
            "   archivé si le fichier temporaire a été nettoyé.",
            "• La dernière tâche vide la zone de transit une fois les fichiers consommés.",
        ],
    )
    slide_contenu(
        prs,
        "KPI & monitoring",
        [
            "• Qualité : validité, association texte-image, images réellement obtenues.",
            "• Volume et diversité : concentration des sources — un jeu capté par un seul",
            "   éditeur transmet son biais au modèle.",
            "• Fraîcheur : âge médian des publications ingérées.",
            "• Performance et coût : durée par étape, quota d'API, disque occupé.",
            "• Les seuils sont définis dans le code : le plan de monitoring et le tableau de",
            "   bord lisent le même dictionnaire et ne peuvent pas diverger.",
        ],
    )
    slide_contenu(
        prs,
        "Résultats d'une exécution réelle",
        [
            "• 179 publications collectées auprès des 4 sources, 134 retenues.",
            "• 100 % d'association texte-image ; 98 % des images demandées obtenues.",
            "• Âge médian des publications : 17 heures.",
            "• Pipeline complet en 60 secondes, 1 appel d'API consommé, 11 Mo d'images.",
            "• DAG exécuté dans Airflow : 5 tâches en succès, DagRun en state=success.",
        ],
    )
    slide_contenu(
        prs,
        "Ce que j'en retiens",
        [
            "• Le plus difficile : rendre les tâches Airflow réellement indépendantes.",
            "   XCom paraissait naturel, mais il enchaînait les tâches.",
            "• L'imprévu : FakeNewsNet, la référence du domaine, ne contient aucune image.",
            "• Deux pièges d'environnement : Docker ne monte pas un dossier Google Drive, et",
            "   installer les dépendances au démarrage cassait la version de SQLAlchemy d'Airflow.",
            "• À surveiller : la concentration des sources, suivie comme un KPI à part entière.",
            "• Suite : déposer le Fakeddit complet, paralléliser l'extraction.",
        ],
    )
    slide_titre(prs, "Merci", "Questions et démonstration en direct du pipeline")

    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    sortie = REPORTS_DIR / "presentation.pptx"
    prs.save(sortie)
    print(f"[ok] présentation enregistrée : {sortie} ({len(prs.slides)} diapositives)")


if __name__ == "__main__":
    main()
